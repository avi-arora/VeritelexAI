"""Document text extraction.

PDF pages with an embedded text layer are extracted locally (fast, free,
exact). Only pages without usable text — scans, images — are sent to Gemini
for OCR, in small page batches that run in parallel. DOCX/XLSX are parsed
with hardened XML parsers (python-docx disables entity resolution; openpyxl
uses defusedxml when installed, which it is).

Large files: the caller streams the source to a scratch file (it is never held
whole in memory) and every CPU-bound parse runs on a small dedicated thread
pool, so the event loop — lease heartbeats, other agent steps, API requests —
stays responsive while a 1,000-page bundle is being read. PDFs are read lazily
through an open file handle (given a *path*, pypdf would load the whole file
into memory).
"""

from __future__ import annotations

import asyncio
import functools
import gc
import io
import itertools
import sys
import threading
from collections.abc import Callable
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO, TypeVar

from google.genai import types
from pypdf import PasswordType, PdfReader, PdfWriter
from pypdf.errors import PyPdfError
from pypdf.generic import StreamObject

from app.agents.schemas import OcrResult, Segment
from app.config import Settings
from app.harness.llm import LLM, LLMResult

OCR_SYSTEM = (
    "You are a meticulous court-records transcriptionist. Transcribe every page of the attached "
    "document verbatim, preserving paragraph numbers, headings, dates, amounts and names exactly. "
    "Do not summarise, translate or correct. Number pages from 1 in the order they appear. "
    "If a page is blank or illegible set readable=false and text to ''."
)
TXT_PAGE_CHARS = 3000
EVICT_STREAM_BYTES = 64 * 1024
"""Streams at least this large (page images, heavy vector content) leave pypdf's cache after use."""
GIL_SWITCH_INTERVAL_S = 0.001
R = TypeVar("R")


@dataclass
class Extracted:
    pages: int
    segments: list[Segment]
    unreadable: list[str] = field(default_factory=list)
    ocr_pages: int = 0
    llm_results: list[LLMResult] = field(default_factory=list)


async def extract(path: Path, ext: str, llm: LLM, s: Settings) -> Extracted:
    """Extract citable text segments from the file at ``path``.

    Raises ``ValueError`` for unsupported or unreadable (corrupt) files.
    """
    if ext == "pdf":
        return await _pdf(path, llm, s)
    if ext == "docx":
        return await _parse(s, _docx, path)
    if ext == "xlsx":
        return await _parse(s, _xlsx, path)
    if ext == "txt":
        return await _parse(s, _txt, path)
    if ext in ("png", "jpg", "tif"):
        mime = {"png": "image/png", "jpg": "image/jpeg", "tif": "image/tiff"}[ext]
        return await _image(await _parse(s, path.read_bytes), mime, llm)
    raise ValueError(f"unsupported type {ext}")


# ------------------------------------------------------------------ parse pool
_pool: ThreadPoolExecutor | None = None
_pool_lock = threading.Lock()


def _get_pool(s: Settings) -> ThreadPoolExecutor:
    global _pool
    with _pool_lock:
        if _pool is None:
            # Parsing is pure Python and holds the GIL. At the default 5 ms switch interval the event
            # loop measured stalls of up to ~900 ms while a 1,000-page PDF was parsed in this pool;
            # at 1 ms the worst stall was ~6 ms. The interval only matters while threads contend
            # for the GIL, so it costs nothing the rest of the time.
            if sys.getswitchinterval() > GIL_SWITCH_INTERVAL_S:
                sys.setswitchinterval(GIL_SWITCH_INTERVAL_S)
            _pool = ThreadPoolExecutor(max_workers=max(1, s.parse_workers), thread_name_prefix="vtx-parse")
        return _pool


async def _parse(s: Settings, fn: Callable[..., R], *args: Any) -> R:
    """Run CPU-bound ``fn`` on the parse pool, kept separate from the default executor that
    serves blocking I/O (GCS), so long parses can never starve storage calls."""
    return await asyncio.get_running_loop().run_in_executor(_get_pool(s), functools.partial(fn, *args))


# ------------------------------------------------------------------ PDF
@dataclass
class _Scan:
    pages: int
    texts: dict[int, str]
    needs_ocr: list[int]
    protected: bool = False


class _PdfSource:
    """pypdf over an open file handle, so PDF objects are read from disk on demand.

    pypdf objects are not thread-safe: every method holds ``_lock`` and runs on the parse pool,
    never on the event loop.
    """

    def __init__(self, path: Path):
        self._lock = threading.Lock()
        self._fh: BinaryIO = open(path, "rb")  # noqa: SIM115 - closed by close()
        self._cache_seen = 0
        try:
            self.reader = PdfReader(self._fh)
        except Exception as exc:  # noqa: BLE001 - any parser failure means the file is not a usable PDF
            self._fh.close()
            raise ValueError(f"not a readable PDF: {exc}"[:300]) from exc

    def scan(self, min_chars: int) -> _Scan:
        """Text layer of every page; pages with too little text are queued for OCR."""
        with self._lock:
            r = self.reader
            try:
                if r.is_encrypted:
                    try:
                        ok = r.decrypt("")
                    except Exception:  # noqa: BLE001 - unsupported encryption scheme
                        ok = PasswordType.NOT_DECRYPTED
                    # decrypt() reports a wrong password by return value, not by raising.
                    if ok == PasswordType.NOT_DECRYPTED:
                        return _Scan(pages=0, texts={}, needs_ocr=[], protected=True)
                texts: dict[int, str] = {}
                needs_ocr: list[int] = []
                for i, page in enumerate(r.pages, start=1):
                    try:
                        t = (page.extract_text() or "").strip()
                    except Exception:  # noqa: BLE001 - malformed content stream on one page
                        t = ""
                    self._evict_large_streams()
                    if len(t) >= min_chars:
                        texts[i] = t
                    else:
                        needs_ocr.append(i)
                return _Scan(pages=len(r.pages), texts=texts, needs_ocr=needs_ocr)
            except PyPdfError as exc:  # broken page tree
                raise ValueError(f"not a readable PDF: {exc}"[:300]) from exc

    def subset(self, pages: list[int]) -> bytes:
        """A new PDF holding just ``pages`` (1-based) — the payload of one OCR call."""
        with self._lock:
            if self._fh.closed:
                raise RuntimeError("PDF source already closed")
            writer = PdfWriter()
            for p in pages:
                writer.add_page(self.reader.pages[p - 1])
            buf = io.BytesIO()
            writer.write(buf)
            del writer
            self._evict_large_streams()
            # PdfWriter's object graph is cyclic: without a collection it — and the page images it
            # copied — lingers until the next full GC (measured: ~50% of a scanned file retained,
            # 9% with this). One collection per OCR batch is negligible next to the OCR call.
            gc.collect()
            return buf.getvalue()

    def close(self) -> None:
        with self._lock:  # waits for an in-flight parse instead of closing the handle under it
            self._fh.close()

    def _evict_large_streams(self) -> None:
        """Drop big streams from pypdf's object cache once they have been used.

        pypdf keeps every object it resolves for the reader's lifetime — including the raw bytes of
        each scanned page's image, which text extraction loads just to read its ``/Subtype``. Left
        alone, scanning a 500 MB bundle would pull roughly all of it onto the heap. Evicted streams
        are simply re-read from the file if needed again (e.g. for an OCR subset). Only entries added
        since the last call are inspected. Object streams stay cached: re-inflating them is costly.
        """
        cache = getattr(self.reader, "resolved_objects", None)
        if not isinstance(cache, dict):
            return  # pypdf internals changed: degrade to no eviction rather than fail
        for key, obj in list(itertools.islice(cache.items(), self._cache_seen, None)):
            if (isinstance(obj, StreamObject) and obj.get("/Type") != "/ObjStm"
                    and len(getattr(obj, "_data", b"") or b"") >= EVICT_STREAM_BYTES):
                del cache[key]
        self._cache_seen = len(cache)


async def _pdf(path: Path, llm: LLM, s: Settings) -> Extracted:
    src = await _parse(s, _PdfSource, path)
    try:
        out = await _pdf_text(src, llm, s)
    except BaseException:
        # Possibly cancelled (step timeout): hand the close to the pool without awaiting it.
        _get_pool(s).submit(src.close)
        raise
    await _parse(s, src.close)
    return out


async def _pdf_text(src: _PdfSource, llm: LLM, s: Settings) -> Extracted:
    scan = await _parse(s, src.scan, s.min_text_chars_per_page)
    if scan.protected:
        return Extracted(pages=0, segments=[], unreadable=["whole document (password protected)"])
    out = Extracted(pages=scan.pages, segments=[])
    texts = scan.texts
    if scan.needs_ocr:
        await _ocr(src, scan.needs_ocr, texts, out, llm, s)
    out.segments = [Segment(ref=f"p. {p}", text=texts[p]) for p in sorted(texts)]
    return out


async def _ocr(src: _PdfSource, pages: list[int], texts: dict[int, str], out: Extracted, llm: LLM, s: Settings) -> None:
    size = max(1, s.ocr_pages_per_call)
    batches = [pages[i : i + size] for i in range(0, len(pages), size)]
    gate = asyncio.Semaphore(max(1, s.ocr_max_inflight_batches))

    async def one(batch: list[int]) -> tuple[dict[int, str], LLMResult]:
        async with gate:  # bounds how many page subsets are held in memory at once
            pdf = await _parse(s, src.subset, batch)
            return await _ocr_call(pdf, batch, llm)

    results = await asyncio.gather(*(one(b) for b in batches), return_exceptions=True)
    for batch, res in zip(batches, results, strict=True):
        if isinstance(res, BaseException):
            out.unreadable += [f"p. {p}" for p in batch]
            continue
        mapping, llm_res = res
        out.llm_results.append(llm_res)
        for p in batch:
            txt = mapping.get(p)
            if txt:
                texts[p] = txt
                out.ocr_pages += 1
            else:
                out.unreadable.append(f"p. {p}")


async def _ocr_call(pdf: bytes, pages: list[int], llm: LLM) -> tuple[dict[int, str], LLMResult]:
    res = await llm.generate(
        schema=OcrResult,
        system=OCR_SYSTEM,
        contents=[types.Part.from_bytes(data=pdf, mime_type="application/pdf"),
                  f"This excerpt has {len(pages)} pages. Transcribe all of them."],
        tier="flash",
        temperature=0.0,
        max_output_tokens=65_536,
    )
    mapping: dict[int, str] = {}
    for pg in res.data.pages:
        if 1 <= pg.page <= len(pages) and pg.readable and pg.text.strip():
            mapping[pages[pg.page - 1]] = pg.text.strip()
    return mapping, res


# ------------------------------------------------------------------ other formats
async def _image(data: bytes, mime: str, llm: LLM) -> Extracted:
    res = await llm.generate(
        schema=OcrResult, system=OCR_SYSTEM, tier="flash", temperature=0.0,
        contents=[types.Part.from_bytes(data=data, mime_type=mime), "Transcribe this page."],
    )
    text = "\n".join(p.text for p in res.data.pages if p.readable).strip()
    out = Extracted(pages=1, segments=[Segment(ref="p. 1", text=text)] if text else [], llm_results=[res])
    if not text:
        out.unreadable = ["p. 1"]
    else:
        out.ocr_pages = 1
    return out


def _docx(path: Path) -> Extracted:
    from docx import Document  # python-docx parses with resolve_entities=False

    doc = Document(str(path))
    segs: list[Segment] = []
    n = 0
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            n += 1
            segs.append(Segment(ref=f"¶ {n}", text=t))
    for ti, table in enumerate(doc.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        body = "\n".join(r for r in rows if r.strip(" |"))
        if body:
            segs.append(Segment(ref=f"Table {ti}", text=body))
    pages = max(1, sum(len(s.text) for s in segs) // 3000)
    return Extracted(pages=pages, segments=segs)


def _xlsx(path: Path) -> Extracted:
    from openpyxl import load_workbook  # uses defusedxml when available

    wb = load_workbook(path, read_only=True, data_only=True)  # read_only streams rows from the zip
    segs: list[Segment] = []
    try:
        for ws in wb.worksheets:
            rows: list[str] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 2000:
                    rows.append("… (truncated)")
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    rows.append("\t".join(cells))
            if rows:
                segs.append(Segment(ref=f"Sheet {ws.title}", text="\n".join(rows)))
    finally:
        wb.close()
    return Extracted(pages=max(1, len(segs)), segments=segs)


def _txt(path: Path) -> Extracted:
    """Read in page-sized pieces so a large text file is never decoded in one go."""
    segs: list[Segment] = []
    pages = 0
    with open(path, encoding="utf-8-sig", errors="replace") as fh:
        while chunk := fh.read(TXT_PAGE_CHARS):
            pages += 1
            if text := chunk.strip():
                segs.append(Segment(ref=f"p. {pages}", text=text))
    return Extracted(pages=max(1, pages), segments=segs)
